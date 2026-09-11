"""Strict, side-effect-free Excel parsing for personnel imports."""

from __future__ import annotations

from io import BytesIO
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from .models import PersonnelCreateRequest

EXCEL_COLUMNS = {
    "账号": "username",
    "姓名": "display_name",
    "部门编号": "department_id",
    "部门名称": "department_name",
    "员工编号": "employee_id",
    "邮箱": "email",
    "岗位": "position",
    "手机号": "mobile",
    "状态": "account_status",
}
REQUIRED_HEADERS = {"账号", "姓名", "部门编号", "部门名称"}
MAX_IMPORT_ROWS = 5000
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 1000


class PersonnelExcelError(ValueError):
    def __init__(self, errors: list[str]):
        super().__init__("；".join(errors))
        self.errors = errors


def _validate_xlsx_container(content: bytes) -> None:
    """Reject malformed or highly inflated Office archives before parsing."""
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if "[Content_Types].xml" not in archive.namelist():
                raise PersonnelExcelError(["文件不是有效的 .xlsx 工作簿"])
            if len(entries) > MAX_ARCHIVE_ENTRIES:
                raise PersonnelExcelError(["Excel 内部文件数量异常，已拒绝解析"])
            if sum(entry.file_size for entry in entries) > MAX_UNCOMPRESSED_BYTES:
                raise PersonnelExcelError(["Excel 解压后内容不能超过 50MB"])
            _reject_hidden_or_formula_content(archive)
    except (BadZipFile, ElementTree.ParseError, RuntimeError) as exc:
        raise PersonnelExcelError(["文件不是有效的 .xlsx 工作簿"]) from exc


def _reject_hidden_or_formula_content(archive: ZipFile) -> None:
    """Do not let a workbook conceal accounts or derive fields from formulas."""

    workbook_entries = [
        entry for entry in archive.infolist() if entry.filename == "xl/workbook.xml"
    ]
    worksheet_entries = [
        entry
        for entry in archive.infolist()
        if entry.filename.startswith("xl/worksheets/")
        and entry.filename.endswith(".xml")
    ]
    for entry in [*workbook_entries, *worksheet_entries]:
        with archive.open(entry) as stream:
            for _, element in ElementTree.iterparse(stream, events=("end",)):
                tag = element.tag.rsplit("}", 1)[-1]
                hidden = str(element.attrib.get("hidden", "")).strip().lower()
                state = str(element.attrib.get("state", "")).strip().lower()
                if tag in {"row", "col"} and hidden in {"1", "true"}:
                    raise PersonnelExcelError(["Excel 不能包含隐藏行或隐藏列"])
                if tag == "sheet" and state in {"hidden", "veryhidden"}:
                    raise PersonnelExcelError(["Excel 不能包含隐藏工作表"])
                if tag == "f":
                    raise PersonnelExcelError(["Excel 不能包含公式，请先粘贴为纯文本值"])
                element.clear()


def parse_personnel_excel(content: bytes, source: str) -> list[PersonnelCreateRequest]:
    _validate_xlsx_container(content)
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise PersonnelExcelError(["文件不是有效的 .xlsx 工作簿"]) from exc

    try:
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        raw_headers = next(iterator, None)
        if not raw_headers:
            raise PersonnelExcelError(["Excel 不能为空"])

        headers = [str(value or "").strip() for value in raw_headers]
        missing = sorted(REQUIRED_HEADERS - set(headers))
        unknown = [header for header in headers if header and header not in EXCEL_COLUMNS]
        header_errors = []
        if missing:
            header_errors.append(f"缺少必填列：{'、'.join(missing)}")
        if unknown:
            header_errors.append(f"存在不支持的列：{'、'.join(unknown)}")
        if len([header for header in headers if header]) != len(set(filter(None, headers))):
            header_errors.append("列名不能重复")
        if header_errors:
            raise PersonnelExcelError(header_errors)

        items: list[PersonnelCreateRequest] = []
        errors: list[str] = []
        seen_usernames: set[str] = set()
        for excel_row, values in enumerate(iterator, start=2):
            if excel_row > MAX_IMPORT_ROWS + 1:
                errors.append(f"最多允许导入 {MAX_IMPORT_ROWS} 行")
                break
            if not any(value not in (None, "") for value in values):
                continue
            raw = {
                EXCEL_COLUMNS[header]: values[index]
                for index, header in enumerate(headers)
                if header and index < len(values)
            }
            raw["source"] = source
            raw.setdefault("account_status", "active")
            try:
                item = PersonnelCreateRequest.model_validate(raw)
            except Exception as exc:
                details = getattr(exc, "errors", lambda: [])()
                message = ", ".join(
                    str(error.get("msg", "格式错误")).removeprefix("Value error, ")
                    for error in details
                ) or "格式错误"
                errors.append(f"第 {excel_row} 行：{message}")
                continue
            if item.username in seen_usernames:
                errors.append(f"第 {excel_row} 行：账号 {item.username} 在文件中重复")
                continue
            if item.username == "admin":
                errors.append(f"第 {excel_row} 行：admin 账号不可通过导入修改")
                continue
            seen_usernames.add(item.username)
            items.append(item)
    finally:
        workbook.close()
    if errors:
        raise PersonnelExcelError(errors[:100])
    if not items:
        raise PersonnelExcelError(["Excel 中没有可导入的人员记录"])
    return items
